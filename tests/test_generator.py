"""
Unit tests for consulting PPT generator.
"""

import pytest
from pathlib import Path
from pptx import Presentation

from ppt.config.theme import ConsultingTheme, Color
from ppt.config.default_theme import load_theme
from ppt.parser.schema import DeckSpec
from ppt.parser.builder import DeckBuilder


def test_theme_loading():
    theme = load_theme("consulting_blue")
    assert theme.name == "consulting_blue"
    assert theme.dimensions.slide_width_inches == 10.0
    assert theme.dimensions.slide_height_inches == 5.625


def test_yaml_spec_validation():
    yaml_path = Path("examples/sample_deck.yaml")
    assert yaml_path.exists()
    builder = DeckBuilder(yaml_path)
    assert len(builder.spec.slides) == 8
    assert builder.spec.title == "China Enterprise Data Platform Solution Proposal"


def test_full_deck_compilation(tmp_path):
    yaml_path = Path("examples/sample_deck.yaml")
    builder = DeckBuilder(yaml_path)
    out_file = tmp_path / "test_deck.pptx"
    prs = builder.build(out_file)

    assert out_file.exists()
    assert len(prs.slides) == 8
    assert prs.slide_width.inches == 10.0
    assert prs.slide_height.inches == 5.625


def test_python_programmatic_api(tmp_path):
    from ppt.core.canvas import SlideCanvas
    from ppt.components.card_grid import CardItem, render_card_grid

    prs = Presentation()
    theme = load_theme("consulting_blue")
    prs.slide_width = prs.slide_width
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    canvas = SlideCanvas(slide, theme)

    canvas.add_header(
        tracker="TEST TRACKER",
        action_title="Action Title for Test",
        lead_note="Lead note explaining this test slide",
    )
    cards = [
        CardItem(title="Card 1", bullets=["Point A", "Point B"]),
        CardItem(title="Card 2", bullets=["Point C", "Point D"], highlight=True),
    ]
    render_card_grid(canvas, cards)
    canvas.add_takeaway_footer("Final takeaway point.")

    out_file = tmp_path / "prog_test.pptx"
    prs.save(str(out_file))
    assert out_file.exists()
