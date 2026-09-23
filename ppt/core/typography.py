"""
Typography utilities and text formatting for consulting presentations.
"""

from typing import Optional, List, Tuple
from pptx.text.text import _Paragraph, _Run, TextFrame
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from ppt.config.theme import ConsultingTheme


def format_run(
    run: _Run,
    text: str,
    font_name: str,
    font_size_pt: float,
    color: RGBColor,
    bold: bool = False,
    italic: bool = False,
) -> _Run:
    """Format an individual text run."""
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(font_size_pt)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    return run


def add_formatted_paragraph(
    text_frame: TextFrame,
    text: str,
    theme: ConsultingTheme,
    font_size_pt: Optional[float] = None,
    color: Optional[RGBColor] = None,
    bold: bool = False,
    font_name: Optional[str] = None,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    space_after_pt: float = 2.0,
    space_before_pt: float = 0.0,
    prefix_bold_colon: bool = True,
) -> _Paragraph:
    """
    Add a paragraph with automatic lead-in bolding for consulting bullet points.
    E.g., "核心痛点：业务数据分散..." -> "核心痛点：" bolded, rest normal.
    """
    if len(text_frame.paragraphs) == 1 and text_frame.paragraphs[0].text == "":
        p = text_frame.paragraphs[0]
    else:
        p = text_frame.add_paragraph()

    p.alignment = align
    p.space_after = Pt(space_after_pt)
    p.space_before = Pt(space_before_pt)

    actual_font = font_name or theme.fonts.font_body
    actual_size = font_size_pt or theme.fonts.size_body
    actual_color = color or theme.colors.text_primary

    # Check for prefix like "Key: Value" or "关键要点：内容"
    split_char = None
    if prefix_bold_colon and not bold:
        for delim in ["：", ": ", " - "]:
            if delim in text:
                split_char = delim
                break

    if split_char:
        prefix, rest = text.split(split_char, 1)
        r1 = p.add_run()
        format_run(r1, prefix + split_char, actual_font, actual_size, actual_color, bold=True)
        r2 = p.add_run()
        format_run(r2, rest, actual_font, actual_size, actual_color, bold=False)
    else:
        r = p.add_run()
        format_run(r, text, actual_font, actual_size, actual_color, bold=bold)

    return p


def set_text_frame_margins(
    text_frame: TextFrame,
    top: float = 0.05,
    bottom: float = 0.05,
    left: float = 0.08,
    right: float = 0.08,
):
    """Set inner margins of a text frame in inches."""
    from pptx.util import Inches
    text_frame.margin_top = Inches(top)
    text_frame.margin_bottom = Inches(bottom)
    text_frame.margin_left = Inches(left)
    text_frame.margin_right = Inches(right)
    text_frame.word_wrap = True
