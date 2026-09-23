"""
Shape drawing helpers for consulting slides.
"""

from typing import Optional, Tuple
from pptx.slide import Slide
from pptx.shapes.autoshape import Shape
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

from ppt.config.theme import ConsultingTheme


def add_rounded_card(
    slide: Slide,
    left: float,
    top: float,
    width: float,
    height: float,
    fill_color: Optional[RGBColor] = None,
    border_color: Optional[RGBColor] = None,
    border_width_pt: float = 1.0,
    name: Optional[str] = None,
) -> Shape:
    """Add a rounded rectangle container card."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    if name:
        shape.name = name

    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()

    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(border_width_pt)
    else:
        shape.line.fill.background()

    return shape


def add_rect_card(
    slide: Slide,
    left: float,
    top: float,
    width: float,
    height: float,
    fill_color: Optional[RGBColor] = None,
    border_color: Optional[RGBColor] = None,
    border_width_pt: float = 1.0,
    name: Optional[str] = None,
) -> Shape:
    """Add a clean rectangular container."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    if name:
        shape.name = name

    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()

    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(border_width_pt)
    else:
        shape.line.fill.background()

    return shape


def add_badge(
    slide: Slide,
    left: float,
    top: float,
    width: float,
    height: float,
    text: str,
    theme: ConsultingTheme,
    bg_color: Optional[RGBColor] = None,
    text_color: Optional[RGBColor] = None,
    font_size_pt: Optional[float] = None,
    rounded: bool = True,
    name: Optional[str] = None,
) -> Shape:
    """Add a small badge/pill (e.g. '01', 'Phase 1', '● 推荐')."""
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(
        shape_type,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    if name:
        shape.name = name

    bg = bg_color or theme.colors.primary
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg
    shape.line.fill.background()

    tf = shape.text_frame
    tf.word_wrap = False
    tf.margin_top = Inches(0.01)
    tf.margin_bottom = Inches(0.01)
    tf.margin_left = Inches(0.03)
    tf.margin_right = Inches(0.03)

    from ppt.core.typography import format_run
    from pptx.enum.text import PP_ALIGN

    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    format_run(
        r,
        text,
        theme.fonts.font_body,
        font_size_pt or theme.fonts.size_badge,
        text_color or theme.colors.text_white,
        bold=True,
    )
    return shape


def add_divider_line(
    slide: Slide,
    left: float,
    top: float,
    width: float,
    color: RGBColor,
    width_pt: float = 0.75,
) -> Shape:
    """Add a horizontal subtle divider line."""
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Pt(width_pt),
    )
    line.fill.solid()
    line.fill.fore_color.rgb = color
    line.line.fill.background()
    return line
