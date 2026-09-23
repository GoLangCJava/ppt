"""
DeckBuilder: Compiles Declarative DeckSpec (JSON/YAML) into native PPTX presentations.
"""

import json
import yaml
from pathlib import Path
from typing import Union, Dict, Any, Optional
from pptx import Presentation
from pptx.util import Inches

from ppt.config.default_theme import load_theme
from ppt.core.canvas import SlideCanvas
from ppt.parser.schema import (
    DeckSpec,
    CoverSlideSpec,
    SectionSlideSpec,
    CardsSlideSpec,
    ArchSlideSpec,
    CompareSlideSpec,
    RoadmapSlideSpec,
    KPISlideSpec,
)
from ppt.components.card_grid import CardItem, render_card_grid
from ppt.components.architecture_view import ArchitectureLayer, ArchitectureModule, render_architecture_view
from ppt.components.comparison_table import (
    ComparisonOption,
    DimensionRow,
    OptionCell,
    RecommendationTakeaway,
    render_comparison_table,
)
from ppt.components.roadmap_timeline import RoadmapPhase, EffortMetric, render_roadmap_timeline
from ppt.components.kpi_summary import KPICard, render_kpi_grid
from ppt.templates.cover_template import render_cover_slide
from ppt.templates.section_template import render_section_slide


class DeckBuilder:
    """Builds and compiles PPTX files from structured specifications."""

    def __init__(self, spec: Union[DeckSpec, Dict[str, Any], str, Path]):
        if isinstance(spec, (str, Path)):
            self.spec = self._load_file(spec)
        elif isinstance(spec, dict):
            self.spec = DeckSpec.model_validate(spec)
        elif isinstance(spec, DeckSpec):
            self.spec = spec
        else:
            raise ValueError(f"Unsupported spec type: {type(spec)}")

        self.theme = load_theme(self.spec.theme)

    def _load_file(self, path: Union[str, Path]) -> DeckSpec:
        p = Path(path)
        if not p.exists():
            raise FileNotFoundError(f"Spec file not found: {p}")
        text = p.read_text(encoding="utf-8")
        if p.suffix.lower() in [".yaml", ".yml"]:
            data = yaml.safe_load(text)
        else:
            data = json.loads(text)
        return DeckSpec.model_validate(data)

    def build(self, output_path: Optional[Union[str, Path]] = None) -> Presentation:
        """Render slides and optionally save presentation."""
        prs = Presentation()

        # Set 16:9 widescreen dimensions (10.0 x 5.625 inches)
        prs.slide_width = Inches(self.theme.dimensions.slide_width_inches)
        prs.slide_height = Inches(self.theme.dimensions.slide_height_inches)
        blank_layout = prs.slide_layouts[6]

        for s_idx, slide_data in enumerate(self.spec.slides):
            slide = prs.slides.add_slide(blank_layout)
            canvas = SlideCanvas(slide, self.theme)

            # Dispatch slide by type
            if isinstance(slide_data, CoverSlideSpec):
                render_cover_slide(
                    canvas,
                    title=slide_data.title,
                    subtitle=slide_data.subtitle,
                    client_tag=slide_data.client_tag,
                    author=slide_data.author or self.spec.author,
                    date_str=slide_data.date or self.spec.date,
                    dark_mode=slide_data.dark_mode,
                )

            elif isinstance(slide_data, SectionSlideSpec):
                render_section_slide(
                    canvas,
                    section_num=slide_data.section_num,
                    title=slide_data.title,
                    subtitle=slide_data.subtitle,
                    deck_tracker=slide_data.deck_tracker or self.spec.title,
                    dark_mode=slide_data.dark_mode,
                )

            elif isinstance(slide_data, CardsSlideSpec):
                canvas.add_header(
                    tracker=slide_data.tracker,
                    action_title=slide_data.action_title,
                    lead_note=slide_data.lead_note,
                )
                if slide_data.takeaway:
                    canvas.add_takeaway_footer(slide_data.takeaway)
                    body_h = canvas.body_height - 0.45
                else:
                    body_h = canvas.body_height

                card_items = [
                    CardItem(
                        title=c.title,
                        badge=c.badge,
                        subtitle=c.subtitle,
                        bullets=c.bullets,
                        tags=c.tags,
                        highlight=c.highlight,
                    )
                    for c in slide_data.cards
                ]
                render_card_grid(canvas, card_items, top=canvas.body_top, height=body_h)

            elif isinstance(slide_data, ArchSlideSpec):
                canvas.add_header(
                    tracker=slide_data.tracker,
                    action_title=slide_data.action_title,
                    lead_note=slide_data.lead_note,
                )
                if slide_data.takeaway:
                    canvas.add_takeaway_footer(slide_data.takeaway)
                    body_h = canvas.body_height - 0.45
                else:
                    body_h = canvas.body_height

                layers = [
                    ArchitectureLayer(
                        layer_num=l.layer_num,
                        name=l.name,
                        sub_en=l.sub_en,
                        modules=[
                            ArchitectureModule(title=m.title, description=m.description, tags=m.tags)
                            for m in l.modules
                        ],
                        protocol_connector=l.protocol_connector,
                    )
                    for l in slide_data.layers
                ]
                render_architecture_view(canvas, layers, top=canvas.body_top, height=body_h)

            elif isinstance(slide_data, CompareSlideSpec):
                canvas.add_header(
                    tracker=slide_data.tracker,
                    action_title=slide_data.action_title,
                    lead_note=slide_data.lead_note,
                )
                opts = [
                    ComparisonOption(name=o.name, subtitle=o.subtitle, is_recommended=o.is_recommended)
                    for o in slide_data.options
                ]
                rows = [
                    DimensionRow(
                        index=r.index,
                        title=r.title,
                        description=r.description,
                        cells=[
                            OptionCell(score_symbol=c.score_symbol, verdict=c.verdict, bullets=c.bullets)
                            for c in r.cells
                        ],
                    )
                    for r in slide_data.rows
                ]
                takeaway = None
                if slide_data.takeaway:
                    takeaway = RecommendationTakeaway(
                        recommended_title=slide_data.takeaway.recommended_title,
                        recommended_text=slide_data.takeaway.recommended_text,
                        tradeoff_title=slide_data.takeaway.tradeoff_title,
                        tradeoff_text=slide_data.takeaway.tradeoff_text,
                    )

                render_comparison_table(
                    canvas,
                    options=opts,
                    rows=rows,
                    takeaway=takeaway,
                    top=canvas.body_top,
                    height=canvas.body_height,
                )

            elif isinstance(slide_data, RoadmapSlideSpec):
                canvas.add_header(
                    tracker=slide_data.tracker,
                    action_title=slide_data.action_title,
                    lead_note=slide_data.lead_note,
                )
                phases = [
                    RoadmapPhase(
                        phase_id=p.phase_id,
                        duration_weeks=p.duration_weeks,
                        title=p.title,
                        tasks=p.tasks,
                        exit_criteria=p.exit_criteria,
                        is_critical=p.is_critical,
                    )
                    for p in slide_data.phases
                ]
                metrics = [
                    EffortMetric(big_stat=m.big_stat, sub_label=m.sub_label)
                    for m in slide_data.metrics
                ]
                render_roadmap_timeline(
                    canvas,
                    phases=phases,
                    traditional_effort=slide_data.traditional_effort,
                    ai_effort=slide_data.ai_effort,
                    metrics=metrics,
                    summary_text=slide_data.summary_text,
                    top=canvas.body_top,
                    height=canvas.body_height,
                )

            elif isinstance(slide_data, KPISlideSpec):
                canvas.add_header(
                    tracker=slide_data.tracker,
                    action_title=slide_data.action_title,
                    lead_note=slide_data.lead_note,
                )
                if slide_data.takeaway:
                    canvas.add_takeaway_footer(slide_data.takeaway)
                    body_h = canvas.body_height - 0.45
                else:
                    body_h = canvas.body_height

                kpi_cards = [
                    KPICard(
                        value=k.value,
                        label=k.label,
                        description=k.description,
                        tag=k.tag,
                        highlight=k.highlight,
                    )
                    for k in slide_data.kpis
                ]
                render_kpi_grid(canvas, kpi_cards, top=canvas.body_top, height=body_h)

        if output_path:
            out_p = Path(output_path)
            out_p.parent.mkdir(parents=True, exist_ok=True)
            prs.save(str(out_p))

        return prs
