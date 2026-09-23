"""
Section divider / agenda transition slide template.
"""

from typing import Optional, List
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

from ppt.core.canvas import SlideCanvas
from ppt.core.shapes import add_rounded_card, add_rect_card, add_badge
from ppt.core.typography import format_run, set_text_frame_margins


def render_section_slide(
    canvas: SlideCanvas,
    section_num: str,       # e.g. '02' or 'PART 2'
    title: str,             # e.g. 'Technology Solution & Architecture Recommendation'
    subtitle: Optional[str] = None,
    deck_tracker: Optional[str] = None,
    dark_mode: bool = True,
):
    """Render a section divider slide."""
    theme = canvas.theme

    if dark_mode:
        canvas.set_background(theme.colors.bg_dark)
        title_col = theme.colors.text_white
        sub_col = theme.colors.primary_light
        num_bg = theme.colors.primary
        num_text = theme.colors.text_white
    else:
        canvas.set_background(theme.colors.bg_card)
        title_col = theme.colors.text_primary
        sub_col = theme.colors.text_secondary
        num_bg = theme.colors.primary_light
        num_text = theme.colors.primary_dark

    # Tracker at top
    if deck_tracker:
        track_box = canvas.slide.shapes.add_textbox(
            Inches(canvas.margin_left),
            Inches(canvas.margin_top),
            Inches(canvas.content_width),
            Inches(0.24),
        )
        tf_tr = track_box.text_frame
        set_text_frame_margins(tf_tr, top=0, bottom=0, left=0, right=0)
        p_tr = tf_tr.paragraphs[0]
        r_tr = p_tr.add_run()
        format_run(
            r_tr,
            deck_tracker.upper(),
            theme.fonts.font_title,
            theme.fonts.size_caption,
            theme.colors.text_muted,
            bold=False,
        )

    # Big Number Badge
    center_y = 2.0
    add_badge(
        canvas.slide,
        canvas.margin_left,
        center_y,
        width=1.10,
        height=0.45,
        text=section_num,
        theme=theme,
        bg_color=num_bg,
        text_color=num_text,
        font_size_pt=14.0,
        rounded=True,
    )

    # Section Title
    title_box = canvas.slide.shapes.add_textbox(
        Inches(canvas.margin_left),
        Inches(center_y + 0.65),
        Inches(canvas.content_width),
        Inches(1.0),
    )
    tf_t = title_box.text_frame
    set_text_frame_margins(tf_t, top=0, bottom=0, left=0, right=0)
    p_t = tf_t.paragraphs[0]
    r_t = p_t.add_run()
    format_run(
        r_t,
        title,
        theme.fonts.font_title,
        22.0,
        title_col,
        bold=True,
    )

    if subtitle:
        p_sub = tf_t.add_paragraph()
        p_sub.space_before = Pt(6.0)
        r_sub = p_sub.add_run()
        format_run(
            r_sub,
            subtitle,
            theme.fonts.font_body,
            theme.fonts.size_action_title,
            sub_col,
            bold=False,
        )
