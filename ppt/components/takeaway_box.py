"""
Executive takeaway callout card component.
"""

from typing import Optional
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

from ppt.core.canvas import SlideCanvas
from ppt.core.shapes import add_rounded_card
from ppt.core.typography import format_run, set_text_frame_margins


def render_takeaway_box(
    canvas: SlideCanvas,
    text: str,
    title: str = "核心判断与建议",
    top: Optional[float] = None,
    height: float = 0.45,
    bg_color: Optional[RGBColor] = None,
    border_color: Optional[RGBColor] = None,
):
    """Render a standalone executive takeaway banner."""
    theme = canvas.theme
    y = top if top is not None else (canvas.slide_height - canvas.margin_bottom - height)
    bg = bg_color or theme.colors.primary_light
    border = border_color or theme.colors.primary

    card = add_rounded_card(
        canvas.slide,
        canvas.margin_left,
        y,
        canvas.content_width,
        height,
        fill_color=bg,
        border_color=border,
        border_width_pt=1.0,
        name="ExecutiveTakeawayBanner",
    )
    tf = card.text_frame
    set_text_frame_margins(tf, top=0.06, bottom=0.06, left=0.15, right=0.15)
    p = tf.paragraphs[0]

    r_title = p.add_run()
    format_run(
        r_title,
        f"【{title}】 ",
        theme.fonts.font_title,
        theme.fonts.size_body,
        theme.colors.primary_dark,
        bold=True,
    )

    r_text = p.add_run()
    format_run(
        r_text,
        text,
        theme.fonts.font_body,
        theme.fonts.size_body,
        theme.colors.text_primary,
        bold=False,
    )
    return card
