"""
Architecture stack view component (N-Tier Architecture Blueprint).
Faithfully recreates McKinsey/BTS multi-layer system architecture slides.
"""

from dataclasses import dataclass, field
from typing import List, Optional
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from ppt.core.canvas import SlideCanvas
from ppt.core.shapes import add_rounded_card, add_rect_card, add_badge, add_divider_line
from ppt.core.typography import format_run, set_text_frame_margins


@dataclass
class ArchitectureModule:
    """Individual service or component box inside an architecture layer."""
    title: str
    description: str
    tags: List[str] = field(default_factory=list)


@dataclass
class ArchitectureLayer:
    """One horizontal layer in the architecture stack."""
    layer_num: str  # e.g. '01'
    name: str       # e.g. '表现层'
    sub_en: str     # e.g. 'INTERFACE LAYER'
    modules: List[ArchitectureModule]
    protocol_connector: Optional[str] = None  # text on arrow/connector below layer
    bg_color: Optional[RGBColor] = None
    accent_color: Optional[RGBColor] = None


def render_architecture_view(
    canvas: SlideCanvas,
    layers: List[ArchitectureLayer],
    top: Optional[float] = None,
    height: Optional[float] = None,
):
    """
    Render a 3-5 layer enterprise architecture stack.
    """
    num_layers = len(layers)
    if num_layers == 0:
        return

    theme = canvas.theme
    start_y = top if top is not None else canvas.body_top
    avail_h = height if height is not None else canvas.body_height

    # Allocate vertical space for layers and connector strips
    connector_h = 0.16
    total_connector_space = (num_layers - 1) * connector_h
    layer_h = (avail_h - total_connector_space) / num_layers

    label_width = 1.35
    content_x = canvas.margin_left + label_width + 0.12
    content_w = canvas.content_width - label_width - 0.12

    curr_y = start_y
    for l_idx, layer in enumerate(layers):
        accent = layer.accent_color or (
            theme.colors.primary_dark if l_idx == 0 else theme.colors.primary
        )

        # 1. Left Layer Label Box
        label_box = add_rounded_card(
            canvas.slide,
            canvas.margin_left,
            curr_y,
            label_width,
            layer_h,
            fill_color=theme.colors.primary_light if l_idx == 0 else theme.colors.bg_tint,
            border_color=accent,
            border_width_pt=1.0,
            name=f"ArchLayer_{layer.layer_num}_Label",
        )
        tf_label = label_box.text_frame
        set_text_frame_margins(tf_label, top=0.06, bottom=0.06, left=0.08, right=0.08)

        # Number badge inside label
        p_num = tf_label.paragraphs[0]
        p_num.alignment = PP_ALIGN.LEFT
        r_num = p_num.add_run()
        format_run(
            r_num,
            f"{layer.layer_num}  ",
            theme.fonts.font_title,
            theme.fonts.size_badge,
            accent,
            bold=True,
        )

        # Title
        p_name = tf_label.add_paragraph()
        p_name.space_after = Pt(1.0)
        p_name.space_before = Pt(2.0)
        r_name = p_name.add_run()
        format_run(
            r_name,
            layer.name,
            theme.fonts.font_title,
            theme.fonts.size_card_title,
            theme.colors.text_primary,
            bold=True,
        )

        # Subtitle EN
        p_en = tf_label.add_paragraph()
        r_en = p_en.add_run()
        format_run(
            r_en,
            layer.sub_en.upper(),
            theme.fonts.font_body,
            6.5,
            theme.colors.text_muted,
            bold=False,
        )

        # 2. Right Layer Modules Grid
        mod_count = len(layer.modules)
        if mod_count > 0:
            mod_gap = 0.10
            mod_w = (content_w - (mod_gap * (mod_count - 1))) / mod_count

            for m_idx, mod in enumerate(layer.modules):
                m_x = content_x + m_idx * (mod_w + mod_gap)
                mod_box = add_rounded_card(
                    canvas.slide,
                    m_x,
                    curr_y,
                    mod_w,
                    layer_h,
                    fill_color=theme.colors.bg_card,
                    border_color=theme.colors.border_subtle,
                    border_width_pt=0.75,
                    name=f"ArchModule_{layer.layer_num}_{m_idx+1}",
                )
                tf_mod = mod_box.text_frame
                set_text_frame_margins(tf_mod, top=0.06, bottom=0.06, left=0.08, right=0.08)

                # Module Title
                p_mt = tf_mod.paragraphs[0]
                r_mt = p_mt.add_run()
                format_run(
                    r_mt,
                    mod.title,
                    theme.fonts.font_title,
                    theme.fonts.size_body + 0.5,
                    theme.colors.primary_dark,
                    bold=True,
                )

                # Module Description
                p_md = tf_mod.add_paragraph()
                p_md.space_before = Pt(2.0)
                r_md = p_md.add_run()
                format_run(
                    r_md,
                    mod.description,
                    theme.fonts.font_body,
                    theme.fonts.size_caption,
                    theme.colors.text_secondary,
                    bold=False,
                )

        curr_y += layer_h

        # 3. Connector strip between layers
        if l_idx < num_layers - 1:
            conn_text = layer.protocol_connector or "↕  集成与数据交互链路"
            conn_box = canvas.slide.shapes.add_textbox(
                Inches(content_x),
                Inches(curr_y),
                Inches(content_w),
                Inches(connector_h),
            )
            tf_conn = conn_box.text_frame
            set_text_frame_margins(tf_conn, top=0, bottom=0, left=0, right=0)
            p_conn = tf_conn.paragraphs[0]
            p_conn.alignment = PP_ALIGN.CENTER
            r_conn = p_conn.add_run()
            format_run(
                r_conn,
                f"──  {conn_text}  ──",
                theme.fonts.font_mono,
                6.5,
                theme.colors.text_muted,
                bold=False,
            )
            curr_y += connector_h
