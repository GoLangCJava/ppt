"""
Multi-dimensional evaluation matrix / comparison table component.
Faithfully recreates Slide 26 (Architecture Option Evaluation & Recommendation).
"""

from dataclasses import dataclass, field
from typing import List, Optional
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from ppt.core.canvas import SlideCanvas
from ppt.core.shapes import add_rounded_card, add_rect_card, add_badge
from ppt.core.typography import format_run, set_text_frame_margins


@dataclass
class OptionCell:
    """Rating and details for an option under one evaluation dimension."""
    score_symbol: str  # '●', '▲', '○'
    verdict: str       # e.g. '优势明显', '存在明显局限'
    bullets: List[str] = field(default_factory=list)


@dataclass
class DimensionRow:
    """One evaluation dimension across multiple options."""
    index: str         # '01', '02', ...
    title: str         # '业务能力与数据可用性'
    description: str   # '平台是否能够完整承接历史数据...'
    cells: List[OptionCell] = field(default_factory=list)


@dataclass
class ComparisonOption:
    """Column definition for candidate solution."""
    name: str          # '方案 A', '方案 B2'
    subtitle: str      # 'PG + Appsmith', '湖仓 + Databricks'
    is_recommended: bool = False


@dataclass
class RecommendationTakeaway:
    """Bottom summary cards."""
    recommended_title: str = "推荐方案"
    recommended_text: str = ""
    tradeoff_title: str = "关键取舍 (Trade-off)"
    tradeoff_text: str = ""


def render_comparison_table(
    canvas: SlideCanvas,
    options: List[ComparisonOption],
    rows: List[DimensionRow],
    takeaway: Optional[RecommendationTakeaway] = None,
    top: Optional[float] = None,
    height: Optional[float] = None,
):
    """
    Render a consulting comparison matrix with Harvey-ball ratings and highlighted recommended column.
    """
    theme = canvas.theme
    start_y = top if top is not None else canvas.body_top
    avail_h = height if height is not None else canvas.body_height

    # Reserve space for bottom takeaway card if present
    takeaway_h = 0.65 if takeaway else 0.0
    table_h = avail_h - takeaway_h - (0.10 if takeaway else 0.0)

    header_h = 0.32
    num_rows = len(rows)
    row_h = (table_h - header_h) / max(num_rows, 1)

    # Column widths: Left dimension column = 1.60 in, rest split among options
    dim_w = 1.60
    rest_w = canvas.content_width - dim_w
    num_opts = len(options)
    opt_w = rest_w / max(num_opts, 1)

    # 1. Render Table Column Headers
    curr_x = canvas.margin_left + dim_w
    for opt_idx, opt in enumerate(options):
        is_rec = opt.is_recommended
        header_bg = theme.colors.primary if is_rec else theme.colors.bg_tint
        text_col = theme.colors.text_white if is_rec else theme.colors.text_primary

        h_card = add_rounded_card(
            canvas.slide,
            curr_x,
            start_y,
            opt_w - 0.04,
            header_h,
            fill_color=header_bg,
            border_color=theme.colors.primary if is_rec else theme.colors.border_subtle,
            border_width_pt=1.0,
            name=f"CompHeader_{opt_idx+1}",
        )
        tf_h = h_card.text_frame
        set_text_frame_margins(tf_h, top=0.04, bottom=0.04, left=0.08, right=0.08)
        p_h = tf_h.paragraphs[0]
        p_h.alignment = PP_ALIGN.CENTER
        
        r_name = p_h.add_run()
        prefix = "★ " if is_rec else ""
        format_run(
            r_name,
            f"{prefix}{opt.name} ｜ {opt.subtitle}",
            theme.fonts.font_title,
            theme.fonts.size_body + 0.5,
            text_col,
            bold=True,
        )
        curr_x += opt_w

    # 2. Render Table Rows
    curr_y = start_y + header_h + 0.04
    for r_idx, row in enumerate(rows):
        # Draw dimension box on the left
        dim_card = add_rounded_card(
            canvas.slide,
            canvas.margin_left,
            curr_y,
            dim_w - 0.05,
            row_h - 0.04,
            fill_color=theme.colors.bg_tint,
            border_color=theme.colors.border_subtle,
            border_width_pt=0.75,
            name=f"Dimension_{row.index}",
        )
        tf_d = dim_card.text_frame
        set_text_frame_margins(tf_d, top=0.05, bottom=0.05, left=0.08, right=0.08)

        p_idx = tf_d.paragraphs[0]
        r_idx_text = p_idx.add_run()
        format_run(
            r_idx_text,
            f"{row.index}  {row.title}\n",
            theme.fonts.font_title,
            theme.fonts.size_card_title,
            theme.colors.primary_dark,
            bold=True,
        )

        p_desc = tf_d.add_paragraph()
        r_desc = p_desc.add_run()
        format_run(
            r_desc,
            row.description,
            theme.fonts.font_body,
            6.5,
            theme.colors.text_muted,
            bold=False,
        )

        # Draw Cells for each option
        cell_x = canvas.margin_left + dim_w
        for opt_idx, opt in enumerate(options):
            cell_data = row.cells[opt_idx] if opt_idx < len(row.cells) else None
            is_rec = opt.is_recommended

            cell_bg = theme.colors.primary_light if is_rec else theme.colors.bg_card
            cell_border = theme.colors.primary if is_rec else theme.colors.border_subtle

            cell_card = add_rounded_card(
                canvas.slide,
                cell_x,
                curr_y,
                opt_w - 0.04,
                row_h - 0.04,
                fill_color=cell_bg,
                border_color=cell_border,
                border_width_pt=1.0 if is_rec else 0.5,
                name=f"Cell_{r_idx+1}_{opt_idx+1}",
            )
            tf_c = cell_card.text_frame
            set_text_frame_margins(tf_c, top=0.04, bottom=0.04, left=0.08, right=0.08)

            if cell_data:
                # Score symbol + Verdict
                p_score = tf_c.paragraphs[0]
                r_sym = p_score.add_run()
                sym_color = (
                    theme.colors.success if cell_data.score_symbol == "●"
                    else theme.colors.warning if cell_data.score_symbol == "▲"
                    else theme.colors.danger
                )
                format_run(
                    r_sym,
                    f"{cell_data.score_symbol}  {cell_data.verdict}\n",
                    theme.fonts.font_body,
                    theme.fonts.size_body,
                    sym_color if not is_rec else theme.colors.primary_dark,
                    bold=True,
                )

                # Bullets
                for b in cell_data.bullets:
                    p_b = tf_c.add_paragraph()
                    p_b.space_before = Pt(1.0)
                    r_b = p_b.add_run()
                    format_run(
                        r_b,
                        f"• {b}",
                        theme.fonts.font_body,
                        6.5,
                        theme.colors.text_secondary,
                        bold=False,
                    )

            cell_x += opt_w

        curr_y += row_h

    # 3. Render Bottom Takeaway / Recommendation & Trade-off Cards
    if takeaway:
        curr_y += 0.06
        rec_w = canvas.content_width * 0.62
        trade_w = canvas.content_width * 0.36
        gap_t = canvas.content_width - rec_w - trade_w

        # Recommended Box
        r_box = add_rounded_card(
            canvas.slide,
            canvas.margin_left,
            curr_y,
            rec_w,
            takeaway_h,
            fill_color=theme.colors.primary_light,
            border_color=theme.colors.primary,
            border_width_pt=1.0,
            name="RecommendationCard",
        )
        tf_r = r_box.text_frame
        set_text_frame_margins(tf_r, top=0.05, bottom=0.05, left=0.10, right=0.10)
        p_rt = tf_r.paragraphs[0]
        r_rt_title = p_rt.add_run()
        format_run(
            r_rt_title,
            f"★ {takeaway.recommended_title}：",
            theme.fonts.font_body,
            theme.fonts.size_body,
            theme.colors.primary_dark,
            bold=True,
        )
        r_rt_body = p_rt.add_run()
        format_run(
            r_rt_body,
            takeaway.recommended_text,
            theme.fonts.font_body,
            theme.fonts.size_caption,
            theme.colors.text_primary,
            bold=False,
        )

        # Trade-off Box
        t_box = add_rounded_card(
            canvas.slide,
            canvas.margin_left + rec_w + gap_t,
            curr_y,
            trade_w,
            takeaway_h,
            fill_color=theme.colors.bg_tint,
            border_color=theme.colors.border_subtle,
            border_width_pt=0.75,
            name="TradeoffCard",
        )
        tf_t = t_box.text_frame
        set_text_frame_margins(tf_t, top=0.05, bottom=0.05, left=0.10, right=0.10)
        p_tt = tf_t.paragraphs[0]
        r_tt_title = p_tt.add_run()
        format_run(
            r_tt_title,
            f"⚖ {takeaway.tradeoff_title}：",
            theme.fonts.font_body,
            theme.fonts.size_body,
            theme.colors.text_primary,
            bold=True,
        )
        r_tt_body = p_tt.add_run()
        format_run(
            r_tt_body,
            takeaway.tradeoff_text,
            theme.fonts.font_body,
            theme.fonts.size_caption,
            theme.colors.text_secondary,
            bold=False,
        )
