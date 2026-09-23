"""
KPI Metric Card Grid component for executive summary and quantitative impact.
"""

from dataclasses import dataclass, field
from typing import List, Optional
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from ppt.core.canvas import SlideCanvas
from ppt.core.shapes import add_rounded_card, add_badge
from ppt.core.typography import format_run, set_text_frame_margins


@dataclass
class KPICard:
    """Individual KPI metric display card."""
    value: str               # e.g. '19个', '40%', '320万'
    label: str               # e.g. '已退役业务系统'
    description: str         # e.g. '包含CRM、费控、追溯等核心系统'
    tag: Optional[str] = None  # e.g. '高优先级', '基线'
    color: Optional[RGBColor] = None
    highlight: bool = False


def render_kpi_grid(
    canvas: SlideCanvas,
    kpis: List[KPICard],
    gap: float = 0.15,
    top: Optional[float] = None,
    height: Optional[float] = None,
):
    """Render a horizontal row of 2-5 high-impact KPI stat cards."""
    count = len(kpis)
    if count == 0:
        return

    cols = canvas.compute_columns(count=count, gap=gap, top=top, height=height)
    theme = canvas.theme

    for i, (kpi, (left, y, w, h)) in enumerate(zip(kpis, cols)):
        accent = kpi.color or (theme.colors.primary if kpi.highlight else theme.colors.primary_dark)
        bg = theme.colors.primary_light if kpi.highlight else theme.colors.bg_card

        card = add_rounded_card(
            canvas.slide,
            left,
            y,
            w,
            h,
            fill_color=bg,
            border_color=accent if kpi.highlight else theme.colors.border_subtle,
            border_width_pt=1.2 if kpi.highlight else 0.75,
            name=f"KPICard_{i+1}",
        )
        tf = card.text_frame
        set_text_frame_margins(tf, top=0.10, bottom=0.08, left=0.12, right=0.12)

        # Big Number Value
        p_val = tf.paragraphs[0]
        r_val = p_val.add_run()
        format_run(
            r_val,
            f"{kpi.value}\n",
            theme.fonts.font_title,
            theme.fonts.size_hero,
            accent,
            bold=True,
        )

        # Label Title
        p_label = tf.add_paragraph()
        p_label.space_before = Pt(2.0)
        r_label = p_label.add_run()
        format_run(
            r_label,
            f"{kpi.label}\n",
            theme.fonts.font_title,
            theme.fonts.size_card_title,
            theme.colors.text_primary,
            bold=True,
        )

        # Description
        p_desc = tf.add_paragraph()
        p_desc.space_before = Pt(3.0)
        r_desc = p_desc.add_run()
        format_run(
            r_desc,
            kpi.description,
            theme.fonts.font_body,
            theme.fonts.size_caption,
            theme.colors.text_secondary,
            bold=False,
        )

        # Tag
        if kpi.tag:
            add_badge(
                canvas.slide,
                left + w - 0.90,
                y + 0.10,
                width=0.80,
                height=0.22,
                text=kpi.tag,
                theme=theme,
                bg_color=accent,
                font_size_pt=6.5,
            )
