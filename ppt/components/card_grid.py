"""
Multi-column card grid component.
"""

from dataclasses import dataclass, field
from typing import List, Optional
from pptx.slide import Slide
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from ppt.core.canvas import SlideCanvas
from ppt.core.shapes import add_rounded_card, add_badge
from ppt.core.typography import format_run, add_formatted_paragraph, set_text_frame_margins


@dataclass
class CardItem:
    """Individual card model."""
    title: str
    badge: Optional[str] = None
    subtitle: Optional[str] = None
    bullets: List[str] = field(default_factory=list)
    tags: List[str] = field(default_factory=list)
    highlight: bool = False
    bg_color: Optional[RGBColor] = None
    border_color: Optional[RGBColor] = None


def render_card_grid(
    canvas: SlideCanvas,
    cards: List[CardItem],
    gap: float = 0.15,
    top: Optional[float] = None,
    height: Optional[float] = None,
):
    """
    Render a horizontal row of 2-5 cards with consulting typography and visual hierarchy.
    """
    count = len(cards)
    if count == 0:
        return

    cols = canvas.compute_columns(count=count, gap=gap, top=top, height=height)

    for i, (card, (left, y, w, h)) in enumerate(zip(cards, cols)):
        theme = canvas.theme
        
        # Decide colors
        if card.highlight:
            bg = card.bg_color or theme.colors.primary_light
            border = card.border_color or theme.colors.primary
            border_pt = 1.5
        else:
            bg = card.bg_color or theme.colors.bg_card
            border = card.border_color or theme.colors.border_subtle
            border_pt = 1.0

        # Draw outer card container
        container = add_rounded_card(
            canvas.slide,
            left,
            y,
            w,
            h,
            fill_color=bg,
            border_color=border,
            border_width_pt=border_pt,
            name=f"Card_{i+1}",
        )

        inner_margin = 0.12
        card_content_left = left + inner_margin
        card_content_w = w - (inner_margin * 2)
        curr_y = y + 0.10

        # Draw Badge (e.g. '01' or '阶段一')
        if card.badge:
            badge_color = theme.colors.primary if card.highlight else theme.colors.primary_dark
            add_badge(
                canvas.slide,
                card_content_left,
                curr_y,
                width=min(0.65, card_content_w * 0.4),
                height=0.22,
                text=card.badge,
                theme=theme,
                bg_color=badge_color,
                name=f"Card_{i+1}_Badge",
            )
            curr_y += 0.28

        # Draw Card Title
        title_box = canvas.slide.shapes.add_textbox(
            Inches(card_content_left),
            Inches(curr_y),
            Inches(card_content_w),
            Inches(0.28),
        )
        tf_t = title_box.text_frame
        set_text_frame_margins(tf_t, top=0, bottom=0, left=0, right=0)
        p_t = tf_t.paragraphs[0]
        r_t = p_t.add_run()
        format_run(
            r_t,
            card.title,
            theme.fonts.font_title,
            theme.fonts.size_card_title,
            theme.colors.text_primary,
            bold=True,
        )
        curr_y += 0.30

        # Draw Subtitle
        if card.subtitle:
            sub_box = canvas.slide.shapes.add_textbox(
                Inches(card_content_left),
                Inches(curr_y),
                Inches(card_content_w),
                Inches(0.25),
            )
            tf_s = sub_box.text_frame
            set_text_frame_margins(tf_s, top=0, bottom=0, left=0, right=0)
            p_s = tf_s.paragraphs[0]
            r_s = p_s.add_run()
            format_run(
                r_s,
                card.subtitle,
                theme.fonts.font_body,
                theme.fonts.size_caption,
                theme.colors.text_muted,
                bold=False,
            )
            curr_y += 0.28

        # Draw Bullet Points
        if card.bullets:
            # compute available height for bullets
            bullet_box = canvas.slide.shapes.add_textbox(
                Inches(card_content_left),
                Inches(curr_y),
                Inches(card_content_w),
                Inches(max(0.5, h - (curr_y - y) - 0.45)),
            )
            tf_b = bullet_box.text_frame
            set_text_frame_margins(tf_b, top=0, bottom=0, left=0, right=0)

            for bullet in card.bullets:
                add_formatted_paragraph(
                    tf_b,
                    f"• {bullet}",
                    theme=theme,
                    font_size_pt=theme.fonts.size_body,
                    color=theme.colors.text_secondary,
                    space_after_pt=4.0,
                    prefix_bold_colon=True,
                )

        # Draw bottom tags if any
        if card.tags:
            tag_y = y + h - 0.32
            tag_left = card_content_left
            for tag_idx, tag in enumerate(card.tags[:3]):
                tag_w = min(0.85, (card_content_w / len(card.tags)) - 0.05)
                add_badge(
                    canvas.slide,
                    tag_left,
                    tag_y,
                    width=tag_w,
                    height=0.20,
                    text=tag,
                    theme=theme,
                    bg_color=theme.colors.bg_tint,
                    text_color=theme.colors.primary_dark,
                    font_size_pt=6.5,
                    rounded=True,
                )
                tag_left += tag_w + 0.05
